"""Paso 9 · Salida: set final de preguntas por categoría + Excel con todos los pasos.

Entrada : s8_check_judge/{chosen,candidates_checked}.csv · s6_select/* · s4_focus/* · s3_weight/* · s2_classify/*
Salida  : out/s9_output/
    final_questions.csv   una fila por pregunta: categoría, intent, unidad, producto, necesidad, pregunta, peso, volumen
    quality_summary.csv   control de calidad por categoría (estado, cuota, juez, peso, revisión IA, Gallagher)
    quality_alerts.csv    alertas que revisar (ver lib/quality.py)
    results.xlsx          Resumen · Calidad · Alertas · Preguntas · Revisión IA · Coste · Foco · Ranking intents ·
                          Grupos · Candidatas · Clasificación · Ajustes · Hojas juntadas
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from rufus_queries.lib import quality
from rufus_queries.lib.common import Config, Out, step_args

STEP = "s9_output"
F, FB, FT = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True), Font(name="Arial", size=14, bold=True)
HDR = PatternFill("solid", fgColor="D5FFBA")


def add_sheet(wb, title, df):
    """Añade una hoja con cabecera formateada y anchos de columna."""
    ws = wb.create_sheet(title[:31])
    ws.append(list(df.columns))
    for r in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in r])
    for c in ws[1]:
        c.font, c.fill, c.alignment = FB, HDR, Alignment(wrap_text=True, vertical="top")
    if len(df) <= 20000:  # dar formato celda a celda a cientos de miles de filas tarda minutos y no aporta
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = F
    for i, col in enumerate(df.columns, 1):
        w = int(df[col].astype(str).str.len().quantile(0.9)) if len(df) else 10
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(str(col)), w) + 2, 60)
    ws.freeze_panes = "A2"


def main(argv=None):
    """Paso 9: final_questions.csv y results.xlsx."""
    args = step_args(__doc__.splitlines()[0], argv)
    cfg, out = Config.load(args.config), Out(args.out)
    briefs = pd.DataFrame(out.read_json("s6_select", "briefs.json"))
    ch = out.read_csv("s8_check_judge", "chosen.csv")
    fin = briefs.merge(ch[["category", "id", "question", "structure", "judge"]], on=["category", "id"], how="left")
    fin["top_keywords"] = fin.top_keywords.apply(lambda x: "; ".join(x[:4]) if isinstance(x, list) else x)
    fin = fin[
        [
            "category",
            "id",
            "intent",
            "intent_name",
            "unit",
            "product",
            "need",
            "question",
            "structure",
            "judge",
            "weight_pct_category",
            "volume",
            "top_keywords",
        ]
    ]
    out.write_csv(STEP, "final_questions.csv", fin)

    qsum, qal = quality.report(out.base, cfg.selection)
    out.write_csv(STEP, "quality_summary.csv", qsum)
    out.write_csv(STEP, "quality_alerts.csv", qal)
    rev = quality.review(out.base)
    usage = out.base / "usage.csv"
    usage = pd.read_csv(usage) if usage.exists() else pd.DataFrame()

    units = out.read_csv("s4_focus", "units.csv")
    ranks = out.read_csv("s6_select", "intent_ranking.csv")
    groups = out.read_csv("s6_select", "groups_scored.csv")
    cands = out.read_csv("s8_check_judge", "candidates_checked.csv")
    kw = out.read_csv("s3_weight", "classified_norm.csv")
    fixes = out.read_csv("s3_weight", "fixes.csv")
    merges = out.read_csv("s4_focus", "merges.csv")

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    lines = [
        ("Preguntas para Rufus / Alexa for Shopping", FT),
        (f"Keywords: {cfg['inputs']['keywords']} · periodo del volumen: {cfg['inputs'].get('volume_period', '?')}", F),
        (
            f"{cfg.selection.get('n_intents')} intents × {cfg.selection.get('questions_per_intent')} preguntas · "
            f"G compite como intent: {cfg.selection.get('generic_as_intent')} · "
            f"revisión con IA: {'sí' if cfg.selection.get('review', True) else 'no'}",
            F,
        ),
        (f"Proveedor IA: {cfg.llm.get('provider')} · coste estimado {quality.cost(out.base):.2f} $", F),
        ("", F),
    ]
    status = dict(zip(qsum.category, qsum.status, strict=False))
    for cat, g in fin.groupby("category"):
        r = ranks[(ranks.category == cat) & ranks.chosen.astype(str).isin(["True", "filler"])]
        n_al = (qal.category == cat) & qal.level.isin(["error", "warning"]) if len(qal) else []
        lines += [
            (cat, FB),
            ("  Intents: " + " · ".join(f"{x.intent} {x.intent_name} {x.share_pct}%" for x in r.itertuples()), F),
            (
                f"  {(g.question.fillna('').astype(str).str.strip() != '').sum()} preguntas · "
                f"nota media del juez {pd.to_numeric(g.judge, errors='coerce').mean():.2f} · "
                f"{pd.to_numeric(g.weight_pct_category, errors='coerce').sum():.1f}% de la demanda representada",
                F,
            ),
            (f"  Control de calidad: {status.get(cat, '?')} · {int(sum(n_al))} alertas (hoja Alertas)", F),
            ("", F),
        ]
    for i, (t, f) in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=t).font = f
    ws.column_dimensions["A"].width = 130

    add_sheet(wb, "Calidad", qsum)
    add_sheet(wb, "Alertas", qal)
    add_sheet(wb, "Preguntas", fin)
    rcols = ["category", "round", "action", "intent", "unit", "product", "need", "fix", "reason", "weight_pct_category"]
    add_sheet(
        wb, "Revisión IA", rev[[c for c in rcols if c in rev.columns]] if len(rev) else pd.DataFrame(columns=rcols)
    )
    add_sheet(wb, "Coste", usage)
    add_sheet(wb, "Foco (unidades)", units)
    add_sheet(wb, "Ranking intents", ranks)
    add_sheet(wb, "Grupos", groups.drop(columns=[c for c in ["Index"] if c in groups.columns]))
    add_sheet(wb, "Candidatas", cands)
    add_sheet(wb, "Clasificación", kw.sort_values("volume", ascending=False))
    add_sheet(wb, "Ajustes", fixes)
    add_sheet(wb, "Hojas juntadas", merges)
    p = out.path(STEP, "results.xlsx")
    wb.save(p)
    print(f"[{STEP}] {len(fin)} preguntas en {fin.category.nunique()} categorías → {p}")
    for r in fin.itertuples():
        print(f"         {r.id:<6} {r.question}")


if __name__ == "__main__":
    main()
